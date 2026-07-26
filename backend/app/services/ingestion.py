"""
Multi-file governing document ingestion (PDF, DOCX, Excel, CSV).

Deterministic extraction only — no LLM.
Every chunk carries document metadata for Qdrant / evidence UI.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

import pdfplumber
from docx import Document


class IngestionError(ValueError):
    """Raised when a governing document cannot be ingested."""


SUPPORTED_EXTENSIONS: frozenset[str] = frozenset(
    {".pdf", ".docx", ".xlsx", ".xls", ".csv"}
)

_HEADING_RE = re.compile(r"^(#{1,6}\s+|[A-Z0-9][\w\s\-/]{2,60}:\s*$|§\s*\d)")
_REQ_ID_RE = re.compile(r"\b(REQ[-_]?\d+|HMI[-_]?\d+|SAF[-_]?\d+)\b", re.I)
_REV_RE = re.compile(r"\b(?:Rev(?:ision)?|REV)\s*[:\s]?\s*([A-Z0-9.]+)\b", re.I)


def ingest_governing_documents(
    paths: list[str | Path],
) -> dict[str, Any]:
    """
    Ingest one or more governing documents into unified requirement chunks.

    Returns
    -------
    dict with keys:
      documents, chunk_count, text_blocks (list of requirement chunks)
    """
    all_chunks: list[dict[str, Any]] = []
    documents: list[dict[str, Any]] = []

    for raw in paths:
        path = Path(raw)
        if not path.is_file():
            raise IngestionError(f"Document not found: {path}")
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise IngestionError(
                f"Unsupported document type '{suffix}' for {path.name}. "
                f"Supported: {sorted(SUPPORTED_EXTENSIONS)}"
            )
        if suffix == ".pdf":
            result = _ingest_pdf(path)
        elif suffix == ".docx":
            result = _ingest_docx(path)
        elif suffix == ".csv":
            result = _ingest_csv(path)
        else:
            result = _ingest_excel(path)

        documents.append(
            {
                "document_name": result["document_name"],
                "revision": result.get("revision") or "",
                "format": result["format"],
                "chunk_count": len(result["text_blocks"]),
            }
        )
        all_chunks.extend(result["text_blocks"])

    return {
        "documents": documents,
        "chunk_count": len(all_chunks),
        "block_count": len(all_chunks),
        "text_blocks": all_chunks,
    }


def ingest_single_document(path: str | Path) -> dict[str, Any]:
    """Back-compat wrapper matching parse_document shape."""
    result = ingest_governing_documents([path])
    if not result["documents"]:
        raise IngestionError("No documents ingested")
    meta = result["documents"][0]
    return {
        "source_file": str(path),
        "format": meta["format"],
        "block_count": result["chunk_count"],
        "text_blocks": result["text_blocks"],
        "document_name": meta["document_name"],
        "revision": meta["revision"],
    }


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _ingest_pdf(path: Path) -> dict[str, Any]:
    document_name = path.name
    revision = ""
    chunks: list[dict[str, Any]] = []
    current_section = ""

    try:
        with pdfplumber.open(path) as pdf:
            for page_index, page in enumerate(pdf.pages, start=1):
                raw = page.extract_text() or ""
                if not revision:
                    m = _REV_RE.search(raw)
                    if m:
                        revision = m.group(1)
                for para_index, paragraph in enumerate(
                    _split_paragraphs(raw), start=1
                ):
                    if _looks_like_heading(paragraph):
                        current_section = paragraph.strip()[:200]
                    req_id = _extract_req_id(paragraph) or ""
                    chunks.append(
                        _chunk(
                            text=paragraph,
                            document_name=document_name,
                            revision=revision,
                            requirement_id=req_id,
                            page=page_index,
                            section=current_section or f"Page {page_index}",
                            source_type="pdf",
                            block_id=f"pdf-p{page_index}-b{para_index}",
                        )
                    )
    except Exception as exc:
        raise IngestionError(f"Failed to parse PDF {path}: {exc}") from exc

    return {
        "document_name": document_name,
        "revision": revision,
        "format": "pdf",
        "text_blocks": chunks,
    }


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def _ingest_docx(path: Path) -> dict[str, Any]:
    document_name = path.name
    revision = ""
    chunks: list[dict[str, Any]] = []
    current_section = ""

    try:
        document = Document(str(path))
    except Exception as exc:
        raise IngestionError(f"Failed to parse DOCX {path}: {exc}") from exc

    # Capture revision from early paragraphs
    for para in document.paragraphs[:30]:
        m = _REV_RE.search(para.text or "")
        if m:
            revision = m.group(1)
            break

    index = 0
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style_name = ""
        try:
            if para.style and para.style.name:
                style_name = str(para.style.name)
        except Exception:
            style_name = ""

        is_heading = style_name.lower().startswith("heading") or _looks_like_heading(
            text
        )
        if is_heading:
            current_section = text[:200]
            continue

        index += 1
        req_id = _extract_req_id(text) or ""
        chunks.append(
            _chunk(
                text=text,
                document_name=document_name,
                revision=revision,
                requirement_id=req_id,
                page=0,
                section=current_section or style_name or "Body",
                source_type="docx",
                block_id=f"docx-p{index}",
            )
        )

    return {
        "document_name": document_name,
        "revision": revision,
        "format": "docx",
        "text_blocks": chunks,
    }


# ---------------------------------------------------------------------------
# Excel / CSV
# ---------------------------------------------------------------------------


def _ingest_excel(path: Path) -> dict[str, Any]:
    try:
        import openpyxl
    except ImportError as exc:
        raise IngestionError(
            "openpyxl is required for Excel ingestion. Install with: uv add openpyxl"
        ) from exc

    document_name = path.name
    revision = ""
    chunks: list[dict[str, Any]] = []

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise IngestionError(f"Failed to parse Excel {path}: {exc}") from exc

    try:
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [
                str(c).strip() if c is not None else f"col_{i}"
                for i, c in enumerate(rows[0])
            ]
            for row_idx, row in enumerate(rows[1:], start=2):
                cells = list(row)
                if all(c is None or str(c).strip() == "" for c in cells):
                    continue
                pairs: list[str] = []
                req_id = ""
                for h, c in zip(headers, cells, strict=False):
                    if c is None or str(c).strip() == "":
                        continue
                    val = str(c).strip()
                    pairs.append(f"{h}: {val}")
                    if not req_id:
                        found = _extract_req_id(val) or _extract_req_id(h)
                        if found:
                            req_id = found
                    if not revision:
                        m = _REV_RE.search(val)
                        if m:
                            revision = m.group(1)
                sentence = (
                    f"On sheet '{sheet.title}' row {row_idx}, "
                    + "; ".join(pairs)
                    + "."
                )
                chunks.append(
                    _chunk(
                        text=sentence,
                        document_name=document_name,
                        revision=revision,
                        requirement_id=req_id or f"{sheet.title}-R{row_idx}",
                        page=0,
                        section=f"Sheet:{sheet.title}",
                        source_type="xlsx",
                        block_id=f"xlsx-{sheet.title}-r{row_idx}",
                        extra={"sheet": sheet.title, "row": row_idx},
                    )
                )
    finally:
        wb.close()

    return {
        "document_name": document_name,
        "revision": revision,
        "format": "xlsx",
        "text_blocks": chunks,
    }


def _ingest_csv(path: Path) -> dict[str, Any]:
    document_name = path.name
    revision = ""
    chunks: list[dict[str, Any]] = []

    try:
        with path.open("r", encoding="utf-8-sig", newline="") as fh:
            reader = csv.DictReader(fh)
            if not reader.fieldnames:
                raise IngestionError(f"CSV has no header row: {path}")
            for row_idx, row in enumerate(reader, start=2):
                pairs = [
                    f"{k}: {v}"
                    for k, v in row.items()
                    if k and v is not None and str(v).strip()
                ]
                if not pairs:
                    continue
                blob = " ".join(pairs)
                req_id = _extract_req_id(blob) or f"CSV-R{row_idx}"
                if not revision:
                    m = _REV_RE.search(blob)
                    if m:
                        revision = m.group(1)
                sentence = f"In CSV row {row_idx}, " + "; ".join(pairs) + "."
                chunks.append(
                    _chunk(
                        text=sentence,
                        document_name=document_name,
                        revision=revision,
                        requirement_id=req_id,
                        page=0,
                        section="CSV",
                        source_type="csv",
                        block_id=f"csv-r{row_idx}",
                        extra={"row": row_idx},
                    )
                )
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError(f"Failed to parse CSV {path}: {exc}") from exc

    return {
        "document_name": document_name,
        "revision": revision,
        "format": "csv",
        "text_blocks": chunks,
    }


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _chunk(
    *,
    text: str,
    document_name: str,
    revision: str,
    requirement_id: str,
    page: int,
    section: str,
    source_type: str,
    block_id: str,
    extra: dict[str, Any] | None = None,
) -> dict[str, Any]:
    payload: dict[str, Any] = {
        "block_id": block_id,
        "source_type": source_type,
        "text": text.strip(),
        # Canonical metadata for Qdrant / evidence UI
        "document_name": document_name,
        "revision": revision or "",
        "requirement_id": requirement_id or "",
        "page": page,
        "section": section or "",
        # Back-compat keys used by pipeline mapping
        "document": document_name,
        "style": section,
    }
    if extra:
        payload.update(extra)
    return payload


def _split_paragraphs(raw: str) -> list[str]:
    parts = re.split(r"\n\s*\n|\n(?=[A-Z0-9])", raw)
    return [p.strip() for p in parts if p and p.strip() and len(p.strip()) > 8]


def _looks_like_heading(text: str) -> bool:
    t = text.strip()
    if len(t) > 120:
        return False
    if _HEADING_RE.match(t):
        return True
    # Short all-caps lines often section titles
    letters = [c for c in t if c.isalpha()]
    if letters and sum(1 for c in letters if c.isupper()) / len(letters) > 0.7:
        return len(t) < 80
    return False


def _extract_req_id(text: str) -> str | None:
    m = _REQ_ID_RE.search(text or "")
    if not m:
        return None
    return m.group(1).upper().replace("_", "-")
